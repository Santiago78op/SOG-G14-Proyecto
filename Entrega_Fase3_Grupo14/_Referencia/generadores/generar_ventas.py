"""Genera el plan de ventas de la Fase 3 (RutaMoto, Grupo 14).

Salidas, en Entrega_Fase3_Grupo14/3_Archivos_Datos/:
  - historial_ventas_fase3.json   -> lo que lee cargar_historial_ventas.py en la instancia
  - RutaMoto_ventas_Fase3.xlsx     -> el mismo plan legible, más los pedidos por el sitio y el cuadre

Corre con el .venv del repo (necesita openpyxl):
  .venv/Scripts/python.exe Entrega_Fase3_Grupo14/_Referencia/generadores/generar_ventas.py

Es determinista: la misma semilla da siempre el mismo plan. Si se cambia algo, se vuelve a
generar y se vuelve a revisar la hoja «Cuadre».
"""
import json
import random
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[3]
EXCEL_F2 = RAIZ / "Entrega_Fase2_Grupo14/3_Archivos_Datos/RutaMoto_datos_maestros_Fase2.xlsx"
SALIDA = RAIZ / "Entrega_Fase3_Grupo14/3_Archivos_Datos"
SEMILLA = 14

# ---------------------------------------------------------------- datos de partida
# Existencia esperada al cerrar la Fase 2: Excel + compras recibidas (P00001-P00003)
# - cotizaciones entregadas (S00004 y S00005 del plan). La instancia manda: el script de
# carga compara contra qty_available real antes de escribir nada.
AJUSTES_FASE2 = {
    "FRE-001": +44, "LUB-004": +72, "LUB-007": +24, "ELE-007": +30,  # compras recibidas
    "FRE-004": -4, "LUB-001": -6, "LLA-001": -5, "LLA-006": -10,     # cotizaciones 1 y 2
}
EXTRA = {  # creados desde el sistema en la Fase 2 (sin fila en el Excel)
    "LUB-007": ("Lubricante de cadena en aerosol 400 ml", "Lubricantes y filtros", 65, 38, 0, 6, 24),
    "ELE-007": ("Bombillo LED de stop 1157", "Sistema eléctrico e iluminación", 40, 22, 0, 8, 30),
}
RECEPCION_P00004 = {"LLA-002": 16}  # se recibe el 21-10, antes de los pedidos por el sitio


def leer_maestros():
    wb = openpyxl.load_workbook(EXCEL_F2, data_only=True)
    prods = {}
    for r in wb["1_Productos"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            prods[r[0]] = {"nombre": r[1], "categoria": r[2], "precio": r[3], "costo": r[4]}
    for r in wb["2_Existencias_iniciales"].iter_rows(min_row=2, values_only=True):
        if r[0] in prods:
            prods[r[0]]["stock"] = r[2]
    for r in wb["3_Reglas_reabastecimiento"].iter_rows(min_row=2, values_only=True):
        if r[0] in prods:
            prods[r[0]]["minimo"], prods[r[0]]["maximo"] = r[1], r[2]
    clientes = []
    for r in wb["7_Clientes"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            clientes.append({"nombre": r[0], "segmento": r[2], "ciudad": r[6]})
    seccion = None
    for r in wb["NO_IMPORTAR_carga_manual"].iter_rows(values_only=True):
        if r[0] and str(r[0]).startswith(("A.", "B.", "C.")):
            seccion = str(r[0])[0]
            continue
        if not r[0] or r[0] in ("Nombre", "Referencia interna"):
            continue
        if seccion == "B" and str(r[0]).startswith("ACC-"):
            prods[r[0]] = {"nombre": r[1], "categoria": r[2], "precio": r[3], "costo": r[4],
                           "stock": r[5], "minimo": r[6], "maximo": r[7]}
        elif seccion == "C":
            clientes.append({"nombre": r[0], "segmento": r[2], "ciudad": r[6]})
    for ref, (n, c, p, k, s, mn, mx) in EXTRA.items():
        prods[ref] = {"nombre": n, "categoria": c, "precio": p, "costo": k,
                      "stock": s, "minimo": mn, "maximo": mx}
    for ref, delta in AJUSTES_FASE2.items():
        prods[ref]["stock"] += delta
    return prods, clientes


# ---------------------------------------------------------------- canastas por segmento
# (referencia, cantidad mínima, cantidad máxima, peso). Los talleres compran por volumen.
CANASTAS = {
    "Repartidor": [("LUB-001", 1, 2, 9), ("LUB-004", 1, 2, 8), ("ELE-002", 1, 2, 4), ("FRE-001", 1, 1, 6),
                   ("LLA-005", 1, 1, 4), ("LUB-006", 1, 1, 6), ("FRE-002", 1, 1, 3), ("ACC-003", 1, 1, 5),
                   ("LUB-007", 1, 1, 4), ("ELE-007", 1, 2, 3), ("FRE-006", 1, 1, 2), ("CAS-004", 1, 1, 2),
                   ("LLA-004", 1, 1, 2), ("ACC-004", 1, 1, 2)],
    "Motociclista particular": [("CAS-001", 1, 1, 4), ("CAS-004", 1, 1, 5), ("LUB-006", 1, 1, 5),
                                ("ELE-004", 1, 1, 4), ("ACC-003", 1, 1, 4), ("LUB-002", 1, 2, 4),
                                ("ELE-005", 1, 1, 3), ("FRE-001", 1, 1, 3), ("ACC-005", 1, 1, 2),
                                ("CAS-002", 1, 1, 2), ("LLA-001", 1, 1, 2), ("ACC-006", 1, 1, 2)],
    "Taller mecánico": [("LUB-001", 3, 6, 9), ("LUB-004", 3, 6, 8), ("FRE-001", 2, 4, 7), ("FRE-002", 2, 3, 4),
                        ("LUB-005", 2, 3, 3), ("FRE-006", 2, 4, 4), ("LLA-004", 2, 4, 3), ("LLA-006", 3, 5, 3),
                        ("ELE-003", 2, 3, 3), ("FRE-005", 2, 4, 3), ("FRE-003", 1, 2, 2), ("LUB-007", 2, 3, 3), ("ELE-007", 2, 4, 2)],
    "Entusiasta": [("LUB-003", 1, 2, 5), ("ELE-003", 1, 2, 5), ("CAS-002", 1, 1, 3), ("CAS-005", 1, 1, 2),
                   ("ACC-001", 1, 1, 4), ("ACC-002", 1, 1, 2), ("FRE-003", 1, 1, 2), ("ELE-005", 1, 1, 3),
                   ("CAS-006", 1, 1, 3), ("FRE-004", 1, 1, 2), ("CAS-003", 1, 1, 2), ("ELE-001", 1, 2, 3)],
}
LINEAS_POR_SEGMENTO = {"Repartidor": (1, 3), "Motociclista particular": (1, 3),
                       "Taller mecánico": (1, 3), "Entusiasta": (1, 2)}
PEDIDOS_POR_DIA = {0: 4, 1: 3, 2: 3, 3: 3, 4: 4, 5: 2, 6: 0}  # lunes=0 ... domingo cerrado

INICIO, FIN = date(2026, 10, 5), date(2026, 10, 20)
TZ_GT = timedelta(hours=6)  # Guatemala = UTC-6; Odoo guarda en UTC


def plan_historial(prods, clientes, rng):
    # lo que después se vende por el sitio queda apartado: la historia no lo puede consumir
    apartado = {}
    for _, _, _, _, lineas, flujo, _, _ in PEDIDOS_WEB:
        for r, q in lineas:
            if flujo != "E" and r not in RECEPCION_P00004:
                apartado[r] = apartado.get(r, 0) + q
    stock = {r: p["stock"] - apartado.get(r, 0) for r, p in prods.items()}
    # Clientes que compran en la historia: 36 de 40, y 9 de ellos repiten (45 pedidos).
    compradores = rng.sample(clientes, 36)
    repiten = set(c["nombre"] for c in rng.sample(compradores, 9))
    # las primeras compras van repartidas en el período y las recompras, en la segunda mitad
    rng.shuffle(compradores)
    segundas = [c for c in compradores if c["nombre"] in repiten]
    rng.shuffle(segundas)
    ordenada = compradores[:24] + sorted(compradores[24:] + segundas, key=lambda _: rng.random())
    pedidos, n = [], 0
    dia = INICIO
    while dia <= FIN:
        for _ in range(PEDIDOS_POR_DIA[dia.weekday()]):
            if n >= len(ordenada):
                break
            cli = ordenada[n]
            n += 1
            lo, hi = LINEAS_POR_SEGMENTO[cli["segmento"]]
            canasta = CANASTAS[cli["segmento"]]
            elegidas, lineas = set(), []
            for _ in range(rng.randint(lo, hi)):
                opciones = [c for c in canasta if c[0] not in elegidas and stock[c[0]] - c[1] >= 1]
                if not opciones:
                    break
                ref, qmin, qmax, _ = rng.choices(opciones, weights=[c[3] for c in opciones])[0]
                q = min(rng.randint(qmin, qmax), stock[ref] - 1)  # nunca deja un producto en cero
                elegidas.add(ref)
                stock[ref] -= q
                lineas.append({"producto": ref, "cantidad": q})
            hora = datetime.combine(dia, datetime.min.time()) + timedelta(
                hours=rng.randint(8, 18), minutes=rng.choice([0, 7, 15, 22, 30, 41, 45, 52]))
            # entrega: 9 de cada 10 en menos de 24 h; el resto entre 26 y 40 h
            horas = rng.randint(3, 20) if rng.random() < 0.9 else rng.randint(26, 40)
            pedidos.append({"cliente": cli["nombre"], "segmento": cli["segmento"], "lineas": lineas,
                            "fecha_local": hora, "entrega_local": hora + timedelta(hours=horas)})
        dia += timedelta(days=1)
    for i, p in enumerate(pedidos, 1):
        p["ref"] = f"HIST-{i:03d}"
        # los talleres compran a crédito: sus 3 últimas facturas quedan pendientes de pago
    talleres = [p for p in pedidos if p["segmento"] == "Taller mecánico"]
    pendientes = {p["ref"] for p in talleres[-3:]}
    for p in pedidos:
        p["pagada"] = p["ref"] not in pendientes
    return pedidos, {r: s + apartado.get(r, 0) for r, s in stock.items()}


# ---------------------------------------------------------------- pedidos por el sitio (rol 2 y rol 4)
# Guion fijo: cada pedido prueba algo distinto del enunciado. «Nuevo» = se registra en el sitio.
PEDIDOS_WEB = [
    # id, día, quién, ¿nuevo?, líneas, flujo, qué prueba, responsable
    ("WEB-01", "2026-10-21", "Cliente nuevo 1 (correo del rol 2)", True, [("LUB-006", 1), ("ACC-003", 1)], "A", "Registro + carrito + pago + factura automática + portal", "Rol 2"),
    ("WEB-02", "2026-10-21", "Cliente nuevo 2", True, [("CAS-001", 1), ("CAS-004", 1)], "A", "Compra de dos productos de la misma categoría", "Rol 2"),
    ("WEB-03", "2026-10-21", "Kevin Alexander Pérez Tzoc", False, [("LUB-001", 2), ("LUB-004", 2)], "A", "Cliente de la Fase 2 con acceso al portal (recompra)", "Rol 2"),
    ("WEB-04", "2026-10-21", "Cliente nuevo 3", True, [("LLA-002", 1)], "D", "LLA-002 vuelve al catálogo después de recibir P00004", "Rol 2"),
    ("WEB-05", "2026-10-22", "Taller Moto Express Zona 7", False, [("FRE-001", 4), ("FRE-006", 4), ("LUB-001", 6)], "C", "Taller con tarifa B2B: el sitio muestra el precio con descuento", "Rol 2"),
    ("WEB-06", "2026-10-22", "Cliente nuevo 4", True, [("CAS-002", 1)], "A", "Ticket alto (Q 650)", "Rol 2"),
    ("WEB-07", "2026-10-22", "Cliente nuevo 5", True, [("ELE-004", 1), ("ELE-005", 1)], "E", "Carrito abandonado: queda como cotización en Ventas", "Rol 2"),
    ("WEB-08", "2026-10-22", "Andrea Sofía Cifuentes Rivas", False, [("ACC-005", 1), ("ACC-006", 1)], "A", "Segunda compra de una motociclista particular", "Rol 2"),
    ("WEB-09", "2026-10-22", "Mónica Isabel Arévalo Soto", False, [("CAS-005", 1), ("CAS-006", 1)], "B", "Lead del formulario -> oportunidad -> cotización por correo -> acepta y paga en el portal", "Rol 4"),
    ("WEB-10", "2026-10-23", "Cliente nuevo 6", True, [("FRE-002", 1), ("LLA-005", 1)], "A", "Compra desde un celular (vista móvil)", "Rol 2"),
    ("WEB-11", "2026-10-23", "Taller Hermanos Cutzal", False, [("LUB-004", 6), ("LUB-007", 3)], "B", "Cotización solicitada desde el sitio por un taller, aceptada en el portal", "Rol 4"),
    ("WEB-12", "2026-10-23", "Cliente nuevo 7", True, [("ELE-003", 1)], "A", "Último pedido antes del corte del BI", "Rol 2"),
]
FLUJOS = {
    "A": "Compra directa: registro o ingreso -> catálogo -> carrito -> pago de prueba -> pedido confirmado -> factura automática -> entrega validada -> portal",
    "B": "Cotización desde el sitio: formulario -> lead en CRM -> oportunidad -> cotización enviada por correo -> el cliente la acepta y paga en el portal",
    "C": "Taller B2B: el cliente con tarifa Talleres B2B ve sus precios en el sitio y compra por volumen",
    "D": "Quiebre resuelto: LLA-002 agotado -> se recibe P00004 -> vuelve a venderse en el sitio",
    "E": "Carrito abandonado: el pedido queda como cotización en Ventas (sirve para medir conversión)",
}


def main():
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    rng = random.Random(SEMILLA)
    prods, clientes = leer_maestros()
    pedidos, stock_hist = plan_historial(prods, clientes, rng)
    stock_ini = {r: p["stock"] for r, p in prods.items()}

    # existencia final: historia -> recepción P00004 -> pedidos por el sitio (menos el carrito abandonado)
    stock_fin = dict(stock_hist)
    for r, q in RECEPCION_P00004.items():
        stock_fin[r] += q
    vendidas_web = {}
    for wid, _, _, _, lineas, flujo, _, _ in PEDIDOS_WEB:
        if flujo == "E":
            continue
        for r, q in lineas:
            stock_fin[r] -= q
            vendidas_web[r] = vendidas_web.get(r, 0) + q
    negativos = [r for r, s in stock_fin.items() if s < 0]
    assert not negativos, f"Plan imposible, quedan en negativo: {negativos}"

    def factor(seg):
        return 0.9 if seg == "Taller mecánico" else 1.0

    montos = [sum(prods[l["producto"]]["precio"] * l["cantidad"] for l in p["lineas"]) * factor(p["segmento"])
              for p in pedidos]
    vendidas = {}
    for p in pedidos:
        for l in p["lineas"]:
            vendidas[l["producto"]] = vendidas.get(l["producto"], 0) + l["cantidad"]
    for r, q in vendidas_web.items():
        vendidas[r] = vendidas.get(r, 0) + q
    top5 = sorted(vendidas.items(), key=lambda x: (-x[1], x[0]))[:5]
    bajo_min = sorted(r for r in prods if stock_fin[r] < prods[r]["minimo"])
    en_cero = sorted(r for r in prods if stock_fin[r] == 0)
    distintos = {p["cliente"] for p in pedidos}
    repiten = {c for c in distintos if sum(1 for p in pedidos if p["cliente"] == c) > 1}
    a_tiempo = sum(1 for p in pedidos if (p["entrega_local"] - p["fecha_local"]) < timedelta(hours=24))

    # ---------------------------------------------------------------- JSON para la instancia
    SALIDA.mkdir(parents=True, exist_ok=True)
    data = {
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "equipo_ventas": "Ventas directas (WhatsApp y teléfono)",
        "nota": "Historia simulada del 05-10 al 20-10-2026. Las fechas van en UTC (hora de Guatemala + 6 h).",
        "pedidos": [{
            "ref": p["ref"], "cliente": p["cliente"], "lineas": p["lineas"], "pagada": p["pagada"],
            "fecha_utc": (p["fecha_local"] + TZ_GT).strftime("%Y-%m-%d %H:%M:%S"),
            "entrega_utc": (p["entrega_local"] + TZ_GT).strftime("%Y-%m-%d %H:%M:%S"),
        } for p in pedidos],
        # unidades que la historia saca de cada producto: el script exige al menos esto + 1 en la instancia
        "unidades_requeridas": {r: stock_ini[r] - stock_hist[r] for r in sorted(prods) if stock_ini[r] != stock_hist[r]},
    }
    (SALIDA / "historial_ventas_fase3.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")

    # ---------------------------------------------------------------- Excel legible
    wb = openpyxl.Workbook()
    neg = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1F3864")

    def hoja(titulo, encabezados, filas, anchos):
        ws = wb.create_sheet(titulo)
        ws.append(encabezados)
        for c in ws[1]:
            c.font, c.fill = neg, fondo
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for f in filas:
            ws.append(f)
        for i, a in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = a
        ws.freeze_panes = "A2"
        return ws

    ws = wb.active
    ws.title = "LEEME"
    for t in [
        "RutaMoto · Fase 3 · Plan de ventas (Grupo 14)",
        "",
        "Este archivo NO se importa. Es el plan legible de lo que se carga y se compra en la Fase 3.",
        "1_Historial: 45 pedidos simulados del 05-10 al 20-10 (canal WhatsApp y teléfono). Los carga el rol 3 con",
        "   _Referencia/scripts_instancia/cargar_historial_ventas.py, que lee historial_ventas_fase3.json.",
        "2_Pedidos_sitio: los 12 pedidos que se hacen POR EL SITIO del 21 al 23-10, cada uno prueba algo del enunciado.",
        "3_Flujos: qué demuestra cada flujo (A a E).",
        "4_Cuadre: existencia antes y después, top 5, productos bajo el mínimo y los KPIs que deberían salir en el BI.",
        "",
        "Por qué hay historia simulada: la gráfica de ventas por día, semana y mes (punto 17) necesita más de tres días",
        "de datos. Se declara así en el informe: las ventas del 05 al 20-10 son simuladas; las del 21 al 23 son reales,",
        "hechas en el sitio por el grupo.",
        "",
        "Montos a precio de lista; a los talleres se les aplica la tarifa Talleres B2B (10 %). Si en la instancia el",
        "descuento es otro, cambian los montos, no las cantidades.",
    ]:
        ws.append([t])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110

    filas = []
    for p, m in zip(pedidos, montos):
        for i, l in enumerate(p["lineas"]):
            pr = prods[l["producto"]]
            filas.append([p["ref"] if i == 0 else "", p["fecha_local"].strftime("%Y-%m-%d %H:%M") if i == 0 else "",
                          p["cliente"] if i == 0 else "", p["segmento"] if i == 0 else "",
                          l["producto"], pr["nombre"], l["cantidad"], pr["precio"],
                          round(pr["precio"] * l["cantidad"] * factor(p["segmento"]), 2),
                          round(m, 2) if i == 0 else "",
                          p["entrega_local"].strftime("%Y-%m-%d %H:%M") if i == 0 else "",
                          ("Sí" if p["pagada"] else "No (crédito)") if i == 0 else ""])
    hoja("1_Historial", ["Pedido", "Fecha (GT)", "Cliente", "Segmento", "Producto", "Nombre", "Cant.", "Precio lista",
                         "Subtotal", "Total pedido", "Entregado (GT)", "Pagada"], filas,
         [10, 17, 34, 22, 10, 38, 7, 11, 11, 12, 17, 13])

    filas = []
    for wid, dia, quien, nuevo, lineas, flujo, prueba, resp in PEDIDOS_WEB:
        for i, (r, q) in enumerate(lineas):
            filas.append([wid if i == 0 else "", dia if i == 0 else "", quien if i == 0 else "",
                          ("Se registra" if nuevo else "Ya existe") if i == 0 else "", r, prods[r]["nombre"], q,
                          flujo if i == 0 else "", prueba if i == 0 else "", resp if i == 0 else ""])
    hoja("2_Pedidos_sitio", ["Pedido", "Día", "Cliente", "Cuenta", "Producto", "Nombre", "Cant.", "Flujo",
                             "Qué prueba", "Quién"], filas, [9, 11, 34, 12, 10, 36, 6, 7, 60, 7])
    hoja("3_Flujos", ["Flujo", "Qué demuestra"], [[k, v] for k, v in FLUJOS.items()], [7, 140])

    filas = []
    for r in sorted(prods):
        pr = prods[r]
        filas.append([r, pr["nombre"], pr["categoria"], stock_ini[r], stock_ini[r] - stock_hist[r],
                      RECEPCION_P00004.get(r, 0), vendidas_web.get(r, 0), stock_fin[r], pr["minimo"],
                      "CERO" if stock_fin[r] == 0 else ("BAJO EL MÍNIMO" if stock_fin[r] < pr["minimo"] else "OK")])
    ws = hoja("4_Cuadre", ["Producto", "Nombre", "Categoría", "Existencia al cerrar Fase 2", "Vendido en historia",
                           "Recibido P00004", "Vendido en el sitio", "Existencia final", "Mínimo", "Estado"],
              filas, [10, 38, 30, 14, 12, 11, 11, 12, 9, 16])
    ws.append([])
    kpis = [
        ("Pedidos simulados (historia)", len(pedidos)),
        ("Venta simulada (Q, con tarifa B2B)", round(sum(montos), 2)),
        ("Ticket promedio de la historia (Q) · meta Q 285", round(sum(montos) / len(montos), 2)),
        ("Clientes distintos en la historia", len(distintos)),
        ("Clientes que repiten · recompra (meta 20 %)", f"{len(repiten)} = {len(repiten) / len(distintos):.1%}"),
        ("Entregas en < 24 h (meta 90 %)", f"{a_tiempo} de {len(pedidos)} = {a_tiempo / len(pedidos):.1%}"),
        ("Pedidos por el sitio (sin contar el carrito abandonado)", len(PEDIDOS_WEB) - 1),
        ("Top 5 por unidades (historia + sitio)", ", ".join(f"{r} ({q})" for r, q in top5)),
        ("Productos bajo el mínimo al final", f"{len(bajo_min)}: {', '.join(bajo_min)}"),
        ("Productos en cero al final · quiebre (meta ≤ 5 %)", f"{len(en_cero)} de {len(prods)} = {len(en_cero) / len(prods):.1%}"),
    ]
    for k, v in kpis:
        ws.append([k, "", "", v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    wb.save(SALIDA / "RutaMoto_ventas_Fase3.xlsx")

    for k, v in kpis:
        print(f"{k}: {v}")


if __name__ == "__main__":
    main()
