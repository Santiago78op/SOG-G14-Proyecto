"""Genera la presentación final de la Fase 3 con el mismo diseño de la de la Fase 2.

  .venv/Scripts/python.exe Entrega_Fase3_Grupo14/_Referencia/generadores/generar_presentacion.py

Salida: Entrega_Fase3_Grupo14/5_Presentacion/Presentacion_Fase3_Grupo14.pptx

Lo que todavía no existe va marcado: las capturas son recuadros punteados «[Captura: ...]» y los números
que salen del BI son «—». Las notas del orador dicen qué decir en cada diapositiva; se vacían antes de
entregar, como en la Fase 2.
"""
import io
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parents[3]
BASE = RAIZ / "Entrega_Fase2_Grupo14/5_Presentacion/Presentacion_Fase2_Grupo14.pptx"
SALIDA = RAIZ / "Entrega_Fase3_Grupo14/5_Presentacion/Presentacion_Fase3_Grupo14.pptx"

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
ORO = RGBColor(0xD7, 0x9B, 0x00)
GRIS = RGBColor(138, 155, 181)
TEXTO = RGBColor(89, 89, 89)
CLARO = RGBColor(221, 228, 238)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FILA = RGBColor(0xF3, 0xF6, 0xFA)
FUENTE = "Segoe UI"
PIE = "RutaMoto · Fase 3 · Grupo 14"


def caja(s, x, y, w, h, texto, tam, color=TEXTO, negrita=False, alin=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, linea in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alin
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = linea
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FUENTE, Pt(tam), negrita, color
    return tb


def rect(s, x, y, w, h, color, forma=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    return r


def diapositiva(prs, titulo, subtitulo, presenta, demo=False, notas=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 0.2, 7.5, AZUL)
    caja(s, 0.6, 0.35, 10.3, 0.9, titulo, 34, AZUL, True)
    caja(s, 0.6, 1.05, 10.5, 0.5, subtitulo, 18, GRIS)
    rect(s, 0.62, 1.6, 1.2, 0.04, ORO)
    if demo:
        b = rect(s, 11.0, 0.45, 1.9, 0.5, ORO, MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = b.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = "DEMO EN VIVO"
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FUENTE, Pt(13), True, AZUL
    caja(s, 0.6, 6.95, 8.0, 0.4, f"Presenta: {presenta}", 12, GRIS)
    caja(s, 10.4, 6.95, 2.7, 0.4, PIE, 11, GRIS, alin=PP_ALIGN.RIGHT)
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    return s


def captura(s, x, y, w, h, que):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = FILA
    r.line.color.rgb = GRIS
    r.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tf = r.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"[Captura: {que}]"
    run.font.name, run.font.size, run.font.italic, run.font.color.rgb = FUENTE, Pt(14), True, GRIS


def tabla(s, x, y, w, encabezados, filas, anchos=None, tam=14):
    t = s.shapes.add_table(len(filas) + 1, len(encabezados), Inches(x), Inches(y), Inches(w),
                           Inches(0.45 * (len(filas) + 1))).table
    anchos = anchos or [w / len(encabezados)] * len(encabezados)
    for i, a in enumerate(anchos):
        t.columns[i].width = Inches(a)
    for i, fila in enumerate([encabezados] + filas):
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = AZUL if i == 0 else (FILA if i % 2 == 0 else BLANCO)
            c.text = str(v)
            p = c.text_frame.paragraphs[0]
            r = p.runs[0] if p.runs else p.add_run()
            r.font.name, r.font.size, r.font.bold = FUENTE, Pt(tam), i == 0
            r.font.color.rgb = BLANCO if i == 0 else TEXTO
    return t


def flujo(s, y, pasos, ancho_total=12.0, x0=0.6):
    """Cadena de pentágono + chevrones, como el ciclo de compras de la Fase 2."""
    w = ancho_total / len(pasos)
    for i, (titulo, detalle) in enumerate(pasos):
        forma = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        r = rect(s, x0 + i * w, y, w + 0.15, 0.9, AZUL if i % 2 == 0 else RGBColor(0x2E, 0x55, 0x86), forma)
        tf = r.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = titulo
        run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FUENTE, Pt(13), True, BLANCO
        caja(s, x0 + i * w + 0.15, y + 1.0, w - 0.1, 0.9, detalle, 12, TEXTO, alin=PP_ALIGN.CENTER)


def main():
    base = Presentation(BASE)
    logo = next(sh for sh in base.slides[0].shapes if sh.shape_type == 13).image.blob

    prs = Presentation(BASE)
    for sid in list(prs.slides._sldIdLst):  # se conserva el archivo (tamaño y tema) y se vacían las diapositivas
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)

    # 1 · Portada
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, AZUL)
    s.shapes.add_picture(io.BytesIO(logo), Inches(0.9), Inches(1.3), Inches(1.3), Inches(1.3))
    caja(s, 0.9, 2.8, 11.0, 1.1, "RutaMoto", 54, BLANCO, True)
    caja(s, 0.9, 3.8, 11.0, 0.6, "Transformación Digital en E-Commerce con ERP, CRM y BI", 22, CLARO)
    rect(s, 0.92, 4.55, 1.4, 0.05, ORO)
    caja(s, 0.9, 4.75, 11.0, 0.9, "Fase 3 · Integración, pruebas y análisis final\nGrupo 14 · Sistemas Organizacionales y "
         "Gerenciales 1 · Sección A · 24 de octubre de 2026", 16, CLARO)
    caja(s, 0.9, 6.2, 11.5, 0.6, "Santiago Julián Barrera Reyes 201905884  ·  Jemima Solmaira Chavajay 201801521  ·  "
         "Alberto Josué Hernández Armas 201903553  ·  Brayan Emanuel García 202300848", 10.5, GRIS)

    # 2 · Qué pedía la fase
    s = diapositiva(prs, "Qué pedía la Fase 3", "Cerrar el ciclo: del clic en el sitio al reporte", "Rol 1",
                    notas="Una frase por criterio. Remarcar que la integración es, según el enunciado, lo que más pesa.")
    tabla(s, 0.6, 2.0, 7.4, ["Criterio", "Puntos"], [
        ["Sitio e-commerce", "15"], ["Integración ERP / CRM / web", "15"], ["Reportes BI", "10"],
        ["Conclusiones con datos", "20"], ["Documentación y entregables", "20"], ["Presentación final", "20"]],
        [5.87, 1.53])
    caja(s, 8.4, 2.0, 4.4, 4.0, "«La integración es el aspecto más valorado y determinante para la calificación final.»\n"
         "— Enunciado, comentarios generales\n\nPor eso esta presentación sigue un solo pedido por todos los módulos.", 17)

    # 3 · La integración
    s = diapositiva(prs, "Una sola base, cero conectores", "Por qué la integración no se programa: se hereda", "Rol 3",
                    notas="La decisión es de la Fase 1: Odoo trae el sitio como aplicación oficial. El sitio no copia "
                          "productos ni existencias: escribe en los mismos registros que ventas, inventario y el CRM.")
    flujo(s, 2.3, [("Sitio", "Registro, carrito y pago"), ("Ventas", "Pedido confirmado"),
                   ("Facturación", "Factura automática"), ("Inventario", "Entrega y existencia"),
                   ("CRM", "Cliente, lead y oportunidad"), ("Portal y BI", "El cliente y la gerencia lo ven")])
    caja(s, 0.6, 4.6, 12.0, 1.8, "•  Un evento del sitio crea el registro en el módulo que corresponde, sin exportar ni copiar.\n"
         "•  El único paso manual es físico: la bodega valida la entrega cuando el paquete sale.\n"
         "•  Elegido en la Fase 1: ERPNext y Dolibarr obligaban a sincronizar con una tienda aparte.", 18)

    # 4 · El sitio
    s = diapositiva(prs, "El sitio", "38 productos, 6 categorías, registro, carrito, pago y factura", "Rol 3", True,
                    "Mostrar la tienda en vivo. Si el túnel falla, usar localhost en la laptop de respaldo.")
    captura(s, 0.6, 2.0, 7.4, 4.6, "tienda con las 6 categorías y los productos con imagen y precio")
    caja(s, 8.4, 2.0, 4.4, 4.6, "•  Registro gratuito de clientes.\n•  Tarifa Talleres B2B visible para los talleres.\n"
         "•  Pago con proveedor en modo de prueba.\n•  Factura automática al confirmarse el pago.\n"
         "•  Sin venta de productos agotados.", 17)

    # 5 · Compra de punta a punta
    s = diapositiva(prs, "Una compra de punta a punta", "Flujo A: el mismo número de pedido en cada módulo", "Roles 2 y 3",
                    True, "Rol 2 compra en vivo; rol 3 abre el mismo pedido en Ventas, Facturación e Inventario y valida "
                          "la entrega. Usar un producto con existencia de sobra, fuera del top 5.")
    flujo(s, 2.2, [("Registro", "Cliente nuevo"), ("Carrito", "Catálogo y cantidades"), ("Pago", "Modo de prueba"),
                   ("Pedido", "Ventas, equipo Sitio web"), ("Factura", "Publicada y pagada"),
                   ("Entrega", "Baja la existencia"), ("Portal", "El cliente lo ve")])
    captura(s, 0.6, 4.3, 12.0, 2.4, "el mismo pedido en Ventas, en la factura y en la entrega")

    # 6 · Del sitio al CRM
    s = diapositiva(prs, "Del sitio al CRM", "Flujo B: de un formulario a una venta aceptada en el portal", "Rol 4", True,
                    "WEB-09 o WEB-11. El correo se ve en Mailpit con el enlace al portal.")
    flujo(s, 2.2, [("Formulario", "Contáctanos"), ("Lead", "Entra solo al CRM"), ("Oportunidad", "Pipeline de la Fase 2"),
                   ("Cotización", "Enviada por correo"), ("Portal", "El cliente acepta y paga")])
    captura(s, 0.6, 4.3, 12.0, 2.4, "lead creado desde el sitio y la cotización aceptada en el portal")

    # 7 · Portal
    s = diapositiva(prs, "Portal del cliente", "Pedidos, entregas y facturas, sin llamar a la tienda", "Rol 2", True)
    captura(s, 0.6, 2.0, 7.4, 4.6, "Mi cuenta: lista de pedidos con su estado y una factura en PDF")
    caja(s, 8.4, 2.0, 4.4, 4.6, "Lo que pide el alcance obligatorio del enunciado: que el cliente vea el estado de sus "
         "pedidos y facturas en el sitio.\n\nAcceso: registro propio o invitación desde Contactos.", 17)

    # 8 · Pruebas
    s = diapositiva(prs, "Doce pedidos, cinco flujos", "Cada pedido del plan prueba algo distinto", "Rol 2",
                    notas="Llenar «Resultado» con lo que pasó. Si algo falló, decirlo: es lo que más se valora.")
    tabla(s, 0.6, 2.0, 12.0, ["Flujo", "Qué prueba", "Pedidos", "Resultado"], [
        ["A", "Compra directa: registro → pago → factura → entrega → portal", "WEB-01 a 03, 06, 08, 10, 12", "—"],
        ["B", "Cotización desde el sitio vía CRM, aceptada en el portal", "WEB-09, WEB-11", "—"],
        ["C", "Taller con tarifa B2B", "WEB-05", "—"],
        ["D", "Quiebre resuelto: LLA-002 vuelve tras recibir P00004", "WEB-04", "—"],
        ["E", "Carrito abandonado queda como cotización", "WEB-07", "—"]], [0.9, 6.4, 3.0, 1.7])

    # 9 · Historia
    s = diapositiva(prs, "Con qué datos se analiza", "Historia simulada y ventas reales, siempre separadas", "Roles 1 y 3",
                    notas="Decirlo antes de que lo pregunten: la historia es simulada y está declarada en el informe.")
    tabla(s, 0.6, 2.0, 7.4, ["", "Historia", "Sitio"], [
        ["Período", "5 al 20 oct", "21 al 23 oct"], ["Pedidos", "45", "11 + 1 carrito"],
        ["Canal en Odoo", "Ventas directas", "Sitio web"], ["Origen", "Script por la API", "El grupo, como clientes"]],
        [2.2, 2.6, 2.6])
    caja(s, 8.4, 2.0, 4.4, 4.0, "Sin la historia, la gráfica por día, semana y mes tendría tres puntos.\n\n"
         "Se cargó por el ORM de Odoo: cada pedido tiene su entrega, su factura y su pago, no solo una fecha.", 17)

    # 10 · BI inventario
    s = diapositiva(prs, "BI · Inventario", "Existencias, categorías y stock bajo, sobre la base en vivo", "Rol 4", True)
    captura(s, 0.6, 2.0, 4.0, 4.6, "barras: existencias por producto")
    captura(s, 4.75, 2.0, 4.0, 4.6, "pastel: inventario por categoría")
    captura(s, 8.9, 2.0, 3.9, 4.6, "tabla: productos bajo el mínimo")

    # 11 · BI ventas
    s = diapositiva(prs, "BI · Ventas", "Evolución por día, semana y mes, y los cinco más vendidos", "Rol 4", True,
                    "Hacer una compra antes y mostrar que aparece en la línea de hoy: es la prueba más fuerte de la integración.")
    captura(s, 0.6, 2.0, 7.4, 4.6, "líneas: ventas por día (cambiar a semana y mes en vivo)")
    captura(s, 8.4, 2.0, 4.4, 4.6, "top 5 de productos por cantidad")

    # 12 · KPIs
    s = diapositiva(prs, "Los indicadores contra la meta", "Metas de la Fase 1, resultados del BI", "Rol 1",
                    notas="Los resultados salen del BI el viernes 23. Decir qué es de la historia simulada y qué del sitio.")
    tabla(s, 0.6, 2.0, 12.0, ["Indicador", "Meta", "Resultado", "¿Cumple?"], [
        ["Ticket promedio", "Q 285", "—", "—"], ["Tasa de recompra", "20 %", "—", "—"],
        ["Entregas en menos de 24 h", "90 %", "—", "—"], ["Quiebre de stock", "≤ 5 %", "—", "—"],
        ["Tasa de conversión del sitio", "2.5 %", "—", "Muestra del prototipo"],
        ["Satisfacción postventa", "≥ 4.3 / 5", "No medido", "Va en escalabilidad"]], [5.0, 2.2, 2.4, 2.4])

    # 13 · Escalabilidad
    s = diapositiva(prs, "De prototipo a empresa", "Lo que falta, por prioridad, con el presupuesto de la Fase 1", "Rol 1",
                    notas="La primera fila condiciona a las demás: sin servidor permanente, el sitio deja de existir "
                          "cuando se apaga la laptop.")
    tabla(s, 0.6, 2.0, 12.0, ["Mejora", "Qué resuelve", "Costo"], [
        ["Servidor en la nube", "El punto único de falla: una laptop", "Q 2,400 al año"],
        ["Dominio rutamoto.gt con HTTPS", "Dirección fija y cobro seguro", "Q 700 al año"],
        ["Pasarela de pago real", "Cobro con tarjeta", "Q 1,500 + comisión"],
        ["Encuesta postventa (Odoo)", "Medir la satisfacción", "Sin licencia"],
        ["Zonas de entrega fuera del área metropolitana", "La expansión del caso", "Sin rubro"]], [4.2, 5.0, 2.8])
    caja(s, 0.6, 5.5, 12.0, 0.8, "Las tres mejoras con costo suman Q 4,600 y ya están dentro de los Q 67,100 de la Fase 1.", 17)

    # 14 · Conclusiones
    s = diapositiva(prs, "Conclusiones", "Lo que dicen los datos", "Rol 1",
                    notas="Se escriben el viernes 23 con los números del BI. Tres como mínimo, cada una con su número.")
    caja(s, 0.6, 2.0, 12.0, 3.5, "1.  —\n2.  —\n3.  —", 20)
    caja(s, 0.6, 5.6, 12.0, 0.8, "Gracias. ¿Preguntas?", 32, AZUL, True)

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    prs.save(SALIDA)
    print("Listo:", SALIDA.relative_to(RAIZ), f"({len(prs.slides)} diapositivas)")


if __name__ == "__main__":
    main()
